"""Загрузка расписания ПКПС: выполняется в GitHub Actions, без токена Telegram."""
import os, re, json, hashlib, urllib.request, time
from datetime import datetime, timezone
from html import unescape
from html.parser import HTMLParser
import xlrd
PAGE='https://pkps-perm.ru/students/raspisanie/fakultet_dizayna_i_servisa_ul_chernyshevskogo_11/'
DAYS=['понедельник','вторник','среда','четверг','пятница','суббота','воскресенье']
GROUP=re.compile(r'^[А-ЯЁA-Z]+-\d{2}-\d{1,2}[КK]?(?:-\d+)?$',re.I)
def key(v):return re.sub(r'\s+','',str(v).upper().replace('–','-').replace('—','-'))
def parse(data):
 w=xlrd.open_workbook(file_contents=data,formatting_info=True);out={}
 for s in w.sheets():
  def cell(r,c):
   v=str(s.cell_value(r,c)).strip()
   if v:return v
   for a,b,x,y in s.merged_cells:
    if a<=r<b and x<=c<y:return str(s.cell_value(a,x)).strip()
   return ''
  def groups(r,c):return [key(v) for v in cell(r,c).splitlines() if GROUP.fullmatch(key(v))]
  for r in range(min(s.nrows,26)):
   for c in range(2,s.ncols):
    if not str(s.cell_value(r,c)).strip():continue
    gs=groups(r,c)
    if not gs:continue
    course=0
    for up in range(r-1,max(-1,r-5),-1):
     m=re.search(r'\b(IV|III|II|I|[1-4])\s*курс',cell(up,c),re.I)
     if m:
      course={'I':1,'II':2,'III':3,'IV':4}.get(m[1].upper(),int(m[1]) if m[1].isdigit() else 0);break
    end=next((y-1 for a,b,x,y in s.merged_cells if a==r and x==c),c)
    lessons={};day=''
    for row in range(r+1,s.nrows):
     if groups(row,c):break
     dv=cell(row,0).lower()
     if dv in DAYS:day=dv
     sl=re.match(r'^\s*(VIII|VII|VI|Vl|IV|III|II|V|I|[1-8])\b',cell(row,1),re.I)
     if not sl or not day:continue
     raw=sl[1].upper();num={'I':1,'II':2,'III':3,'IV':4,'V':5,'VI':6,'VL':6,'VII':7,'VIII':8}.get(raw,int(raw) if raw.isdigit() else 0)
     t=re.search(r'(\d{1,2})[.:](\d{2})',cell(row,1));clock=f'{int(t[1]):02}:{t[2]}' if t else ''
     entries=[]
     for col in range(c,end+1):
      text=cell(row,col)
      if text and not any(x['text']==text for x in entries):entries.append({'day':day,'number':num,'time':clock,'text':text,'subgroup':col-c+1})
     for entry in entries:
      if len(entries)==1:entry['subgroup']=None
      lessons.setdefault(day,[]).append(entry)
    for g in gs:out[g]={'group':g,'course':course,'days':lessons,'present':True}
 if not out:raise ValueError('Группы не распознаны')
 return out
class Links(HTMLParser):
 def __init__(self):super().__init__();self.current=None;self.links=[]
 def handle_starttag(self,tag,attrs):
  if tag=='a':self.current=[dict(attrs).get('href',''),'']
 def handle_data(self,data):
  if self.current is not None:self.current[1]+=data
 def handle_endtag(self,tag):
  if tag=='a' and self.current is not None:self.links.append(self.current);self.current=None

def request(url,body=None,secret=None):
 headers={'User-Agent':'PKPS-Schedule/2.0'}
 if body is not None:headers['Content-Type']='application/json'
 if secret:headers['X-Import-Secret']=secret
 req=urllib.request.Request(url,data=json.dumps(body,ensure_ascii=False).encode() if body is not None else None,headers=headers)
 for attempt in range(3):
  try:
   with urllib.request.urlopen(req,timeout=45) as r:return r.read()
  except Exception:
   if attempt==2:raise
   time.sleep(2*(attempt+1))
def import_design():
 from urllib.parse import urljoin,urlsplit
 worker=os.environ['WORKER_URL'].rstrip('/');secret=os.environ['IMPORT_SECRET']
 if not worker.startswith('https://'):raise ValueError('WORKER_URL должен начинаться с https://')
 manifest=json.loads(request(worker+'/manifest',{},secret));p=Links();p.feed(request(PAGE).decode('utf-8'))
 docs={}
 for href,title in p.links:
  url=urljoin(PAGE,unescape(href));title=' '.join(title.split());date=re.search(r'(\d{2})\.(\d{2})\.(\d{4})',title)
  if urlsplit(url).netloc!=urlsplit(PAGE).netloc or not re.search(r'\.xlsx?$',urlsplit(url).path,re.I):continue
  if re.search(r'полугод|учебных занятий',title,re.I):doc_key='base'
  elif date and re.search(r'изменени|расписан',title,re.I):doc_key=f'{date[3]}-{date[2]}-{date[1]}'
  else:continue
  docs.setdefault(doc_key,{'key':doc_key,'url':url,'title':title})
 if 'base' not in docs:raise ValueError('Основное расписание не найдено')
 selected=['base']+sorted([k for k in docs if k!='base'],reverse=True)[:10]
 for k in selected:
  doc=docs[k];raw=request(doc['url']);digest=hashlib.sha256(raw).hexdigest()
  if manifest.get(k)==digest:print(k,'без изменений');continue
  groups=parse(raw)
  # Сначала сохраняем данные; уведомления выполняются отдельно после импорта.
  response=json.loads(request(worker+'/import',{'key':k,'hash':digest,'data':{'groups':groups,'url':doc['url'],'title':doc['title']}},secret))
  if not response.get('ok'):raise ValueError('Ошибка импорта')
  print(k,'импортировано групп:',len(groups))
 request(worker+'/complete',{'dates':selected[1:]},secret)
 print('Обновление завершено')

FP_PAGE='https://pkps-perm.ru/students/raspisanie/fakultet_predprinimatelstva_ul_permskaya_226/'

def fp_parse(raw, expected_date):
 from io import BytesIO
 import openpyxl, warnings
 with warnings.catch_warnings():
  warnings.simplefilter('ignore', UserWarning)
  book=openpyxl.load_workbook(BytesIO(raw),data_only=True)
 out={}
 for sheet in book:
  if not re.search(r'курс',sheet.title,re.I):continue
  def cell(r,c):
   value=sheet.cell(r,c).value
   if value is not None:return str(value).strip()
   for merged in sheet.merged_cells.ranges:
    if merged.min_row<=r<=merged.max_row and merged.min_col<=c<=merged.max_col:
     return str(sheet.cell(merged.min_row,merged.min_col).value or '').strip()
   return ''
  header=next((r for r in range(1,min(sheet.max_row,30)+1) if sum(bool(GROUP.fullmatch(key(cell(r,c)))) for c in range(3,sheet.max_column+1))>=2),None)
  if not header:continue
  dates=re.findall(r'\b(\d{2})\.(\d{2})\.(\d{4})\b',' '.join(cell(r,c) for r in range(1,header) for c in range(1,sheet.max_column+1)))
  if expected_date not in [f'{y}-{m}-{d}' for d,m,y in dates]:raise ValueError('Дата внутри Excel не совпадает со ссылкой: '+expected_date)
  for c in range(3,sheet.max_column+1):
   name=key(cell(header,c))
   if not GROUP.fullmatch(name):continue
   course=0
   for r in range(header-1,0,-1):
    match=re.search(r'\b([1-4])\s*курс',cell(r,c),re.I)
    if match:course=int(match[1]);break
   if not course:raise ValueError('Не определён курс: '+name)
   lessons={};day='';starts=[]
   for r in range(header+1,sheet.max_row+1):
    value=cell(r,1).lower()
    if value in DAYS:day=value
    match=re.match(r'^\s*(VIII|VII|VI|IV|III|II|V|I|[1-8])\b',str(sheet.cell(r,2).value or ''),re.I)
    if match and day:
     num={'I':1,'II':2,'III':3,'IV':4,'V':5,'VI':6,'VII':7,'VIII':8}.get(match[1].upper(),int(match[1]) if match[1].isdigit() else 0)
     clock=re.search(r'(\d{1,2})[.:](\d{2})',cell(r,2))
     starts.append((r,day,num,f'{int(clock[1]):02}:{clock[2]}' if clock else ''))
   for i,(r,day,num,clock) in enumerate(starts):
    end=starts[i+1][0] if i+1<len(starts) else sheet.max_row+1
    # Предмет и преподаватель/кабинет расположены на отдельных строках.
    parts=[]
    for rr in range(r,end):
     text=cell(rr,c)
     if text and text not in parts:parts.append(text)
    if parts:lessons.setdefault(day,[]).append({'day':day,'number':num,'time':clock,'text':'\n'.join(parts),'subgroup':None})
   out['FP~'+name]={'group':name,'course':course,'faculty':'fp','days':lessons,'present':True}
 if not out:raise ValueError('Группы факультета предпринимательства не найдены')
 return out

def fp_documents(html):
 from urllib.parse import urljoin,urlsplit,quote,unquote
 links=Links();links.feed(html);dated={};pdfs={}
 for href,title in links.links:
  url=quote(unquote(urljoin(FP_PAGE,unescape(href))),safe=':/');title=' '.join(title.split())
  if urlsplit(url).netloc!=urlsplit(FP_PAGE).netloc:continue
  if urlsplit(url).path.lower().endswith('.pdf'):
   nums=[int(v) for v in re.findall(r'\b[1-4]\b',title)]
   for course in nums:pdfs[course]=url
  date=re.search(r'(\d{2})\.(\d{2})\.(\d{4})',title)
  if date and urlsplit(url).path.lower().endswith('.xlsx') and re.search(r'изменени|расписан',title,re.I):
   k=f'{date[3]}-{date[2]}-{date[1]}';dated.setdefault(k,{'url':url,'title':title})
 if set(pdfs)!={1,2,3,4} or not dated:raise ValueError('Неполный список источников предпринимательства')
 return pdfs,dated

def import_enterprise():
 worker=os.environ['WORKER_URL'].rstrip('/');secret=os.environ['IMPORT_SECRET']
 manifest=json.loads(request(worker+'/manifest',{},secret))
 pdfs,docs=fp_documents(request(FP_PAGE).decode('utf-8'))
 selected=sorted(docs,reverse=True)[:10]
 # Свежий Excel содержит полный список групп и их курсов.
 newest=selected[0];raw=request(docs[newest]['url']);roster=fp_parse(raw,newest)
 pdf_hash={}
 for url in set(pdfs.values()):pdf_hash[url]=hashlib.sha256(request(url)).hexdigest()
 groups={}
 for gid,g in roster.items():
  url=pdfs[g['course']]
  groups[gid]={**g,'days':{},'pdf':True,'url':url,'source_hash':pdf_hash[url]}
 payload={'groups':groups,'url':FP_PAGE,'title':'Основное расписание факультета предпринимательства (PDF)'}
 digest=hashlib.sha256(json.dumps(payload,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
 def send(k,h,data):
  if manifest.get(k)==h:print(k,'без изменений');return
  result=json.loads(request(worker+'/import',{'key':k,'hash':h,'data':data},secret))
  if not result.get('ok'):raise ValueError('Ошибка импорта '+k)
  print(k,'импортировано групп:',len(data['groups']))
 send('fp:base',digest,payload)
 for date in selected:
  doc=docs[date];data=raw if date==newest else request(doc['url']);digest=hashlib.sha256(data).hexdigest()
  if manifest.get('fp:'+date)==digest:print('fp:'+date,'без изменений');continue
  parsed=roster if date==newest else fp_parse(data,date)
  send('fp:'+date,digest,{'groups':parsed,'url':doc['url'],'title':doc['title']})
 request(worker+'/complete',{'faculty':'fp','dates':selected},secret)
 print('Факультет предпринимательства обновлён')

def main():
 failures=[]
 for name,job in [('Дизайна и сервиса',import_design),('Предпринимательства',import_enterprise)]:
  try:job()
  except Exception as err:failures.append(name+': '+str(err))
 if failures:raise RuntimeError('; '.join(failures))

if __name__=='__main__':main()
